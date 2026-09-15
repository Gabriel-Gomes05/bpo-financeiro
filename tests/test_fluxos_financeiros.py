import asyncio
import inspect
import json
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from urllib.parse import unquote_plus

import pandas as pd
import pytest
from fastapi import Request
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.models import (
    Atendimento, CentroCusto, ClienteBPO, ContaPagar, PagamentoParcialContaPagar,
    PerfilUsuario, PlanoConta, StatusConciliacao, StatusContaPagar, Usuario,
)
from app.routers import contas_pagar, dashboard, fechamento, gestao, lancamentos, plano_contas
from scripts.generate_import_templates import gerar


def request(path="/", empresa=1):
    headers = [(b"cookie", f"cliente_ativo={empresa}".encode())] if empresa else []
    return Request({"type": "http", "method": "GET", "path": path,
                    "headers": headers, "query_string": b"", "session": {}})


def chamar(funcao, **kwargs):
    params = {}
    for name, param in inspect.signature(funcao).parameters.items():
        if param.default is not inspect.Parameter.empty:
            params[name] = getattr(param.default, "default", param.default)
    params.update(kwargs)
    return asyncio.run(funcao(**params))


@pytest.fixture
def dados(monkeypatch):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    for modulo in (contas_pagar, lancamentos):
        monkeypatch.setattr(modulo, "_log", lambda *args, **kwargs: None)
    with Session(engine) as db:
        usuario = Usuario(id=1, nome="Operador", email="teste@example.test", senha_hash="teste", perfil=PerfilUsuario.coordenador)
        receita = PlanoConta(nome="Consultas", tipo="receita", grupo="Receitas", chave="consultas", codigo="R1", ativo=True)
        despesa = PlanoConta(nome="Aluguel", tipo="despesa", grupo="Despesas", chave="df_aluguel", codigo="D1", ativo=True)
        centro = CentroCusto(cliente_id=1, nome="Clinica", codigo="C1", ativo=True)
        db.add_all([ClienteBPO(id=1, nome="A"), ClienteBPO(id=2, nome="B"), usuario, receita, despesa, centro])
        db.commit()
        yield SimpleNamespace(db=db, usuario=usuario, receita=receita, despesa=despesa, centro=centro)
    engine.dispose()


def conta(dados, **kwargs):
    params = dict(cliente_id=1, descricao="Conta", valor=100, vencimento=date(2026, 2, 10),
                  data_competencia=date(2026, 1, 31), status=StatusContaPagar.agendado)
    params.update(kwargs)
    item = ContaPagar(**params)
    dados.db.add(item)
    dados.db.commit()
    return item


def test_aprovacao_exibe_somente_agendados_da_empresa_e_data(dados):
    esperado = conta(dados)
    for status in StatusContaPagar:
        if status != StatusContaPagar.agendado:
            conta(dados, status=status)
    conta(dados, cliente_id=2)
    conta(dados, vencimento=date(2026, 3, 1))
    assert fechamento._contas_para_aprovacao(dados.db, 1, date(2026, 2, 28)) == [esperado]


def test_dashboard_conta_todos_status_abertos(dados):
    for status in StatusContaPagar:
        conta(dados, status=status)
    conta(dados, cliente_id=2)
    blocos = dashboard._blocos_processo(dados.db, date(2026, 2, 28), [1])
    assert blocos[1]["linha1"].startswith("3 ")


def test_dashboard_foco_e_geral_contam_todos_status_abertos(dados):
    for status in StatusContaPagar:
        conta(dados, status=status)
    conta(dados, cliente_id=2, status=StatusContaPagar.agendado)

    foco = dashboard._dashboard_coordenador(dados.db, date(2026, 2, 28), cliente_ativo_id=1)
    assert foco["total_vencendo"] == 3
    assert len(foco["contas_vencendo"]) == 3

    geral = dashboard._dashboard_coordenador(dados.db, date(2026, 2, 28), cliente_ativo_id=None)
    assert geral["total_vencendo"] == 4


def test_dashboard_funcionario_conta_todos_status_abertos(dados):
    funcionario = Usuario(id=2, nome="Func", email="func@example.test", senha_hash="teste", perfil=PerfilUsuario.funcionario)
    dados.db.add(funcionario)
    dados.db.query(ClienteBPO).filter(ClienteBPO.id == 1).update({"funcionario_id": funcionario.id})
    dados.db.commit()
    for status in StatusContaPagar:
        conta(dados, status=status)

    painel = dashboard._dashboard_funcionario(dados.db, date(2026, 2, 28), funcionario)
    assert painel["total_vencendo"] == 3
    assert painel["resumo_clientes"][0]["vencendo"] == 3


def test_receitas_separam_competencia_de_recebimento(dados):
    receita = Atendimento(condicao_pagamento="avista", cliente_id=1, valor_servico=100, valor_liquido=95,
        data_atendimento=date(2026, 1, 31), data_credito=date(2026, 2, 5),
        status_conciliacao=StatusConciliacao.conciliado)
    pendente = Atendimento(condicao_pagamento="avista", cliente_id=1, valor_servico=50, data_atendimento=date(2026, 1, 10),
        data_credito=date(2026, 2, 5), status_conciliacao=StatusConciliacao.pendente)
    dados.db.add_all([receita, pendente])
    dados.db.commit()
    assert len(gestao._receitas_periodo(dados.db, 1, date(2026, 1, 1), date(2026, 1, 31), "competencia")) == 2
    assert gestao._receitas_periodo(dados.db, 1, date(2026, 1, 1), date(2026, 1, 31), "caixa") == []
    assert gestao._receitas_periodo(dados.db, 1, date(2026, 2, 1), date(2026, 2, 28), "caixa") == [receita]


def test_caixa_pagamentos_parciais_nao_duplica_quitacao(dados):
    item = conta(dados, status=StatusContaPagar.pago, data_pagamento=date(2026, 3, 5))
    dados.db.add_all([
        PagamentoParcialContaPagar(conta_pagar_id=item.id, valor=30, data_pagamento=date(2026, 2, 10)),
        PagamentoParcialContaPagar(conta_pagar_id=item.id, valor=70, data_pagamento=date(2026, 3, 5)),
    ])
    dados.db.commit()
    conta(dados, status=StatusContaPagar.cancelado, valor=900)
    outra_empresa = conta(dados, cliente_id=2)
    dados.db.add(PagamentoParcialContaPagar(conta_pagar_id=outra_empresa.id, valor=90, data_pagamento=date(2026, 2, 10)))
    dados.db.commit()
    for mes, esperado in [(1, 0), (2, 30), (3, 70)]:
        _, valores, _ = gestao._despesas_periodo(dados.db, 1, date(2026, mes, 1), date(2026, mes, 28), "caixa")
        assert sum(valores.values()) == esperado
    _, valores, _ = gestao._despesas_periodo(dados.db, 1, date(2026, 1, 1), date(2026, 1, 31), "competencia")
    assert sum(valores.values()) == 100


def test_caixa_inclui_pagamento_legado_sem_parcelas(dados):
    item = conta(dados, status=StatusContaPagar.pago, data_pagamento=date(2026, 2, 15))
    conta(dados, data_pagamento=date(2026, 2, 15), status=StatusContaPagar.cancelado)
    contas, valores, _ = gestao._despesas_periodo(dados.db, 1, date(2026, 2, 1), date(2026, 2, 28), "caixa")
    assert contas == [item]
    assert valores[item.id] == 100


def test_receita_liquida_zero_nao_vira_valor_bruto(dados):
    dados.db.add(Atendimento(condicao_pagamento="avista", cliente_id=1, valor_servico=100, valor_liquido=0,
        data_atendimento=date(2026, 1, 15), data_credito=date(2026, 1, 20),
        status_conciliacao=StatusConciliacao.conciliado))
    dados.db.commit()
    response = chamar(gestao.gestao_receitas, request=request(), cliente_id=1, mes=1, ano=2026,
        regime="caixa", db=dados.db, usuario=dados.usuario)
    assert response.context["total"] == 0
    assert gestao._carregar_dre_periodo(dados.db, 1, 1, 2026, regime="caixa")["total_real_rec"] == 0


def test_dre_respeita_regime(dados):
    conta(dados)
    dados.db.add(Atendimento(condicao_pagamento="avista", cliente_id=1, valor_servico=200, valor_liquido=200,
        data_atendimento=date(2026, 1, 15), status_conciliacao=StatusConciliacao.pendente))
    dados.db.commit()
    competencia = gestao._carregar_dre_periodo(dados.db, 1, 1, 2026, regime="competencia")
    caixa = gestao._carregar_dre_periodo(dados.db, 1, 1, 2026, regime="caixa")
    assert competencia["total_real_rec"] == 200
    assert competencia["total_real_desp"] == 100
    assert caixa["total_real_rec"] == caixa["total_real_desp"] == 0


def test_coordenador_edita_plano_padrao_sem_mudar_chave(dados):
    chamar(plano_contas.editar_plano_conta, request=request(empresa=None), conta_id=dados.receita.id,
        grupo="Receitas", nome="Consultas e exames", codigo="R10", db=dados.db, usuario=dados.usuario)
    assert dados.receita.nome == "Consultas e exames"
    assert dados.receita.chave == "consultas"
    chamar(plano_contas.editar_plano_conta, request=request(empresa=2), conta_id=dados.receita.id,
        grupo="Receitas", nome="Indevido", codigo="R10", db=dados.db, usuario=dados.usuario)
    assert dados.receita.nome == "Consultas e exames"


def criar_receita(dados, plano_id):
    return chamar(lancamentos.criar_lancamentos, request=request(), cliente_id=1,
        data_atendimento=date(2026, 1, 31), plano_conta_id=[str(plano_id)], valor_servico=["100"],
        forma_pagamento=["dinheiro"], rateios_json=[json.dumps([{"centro_custo_id": dados.centro.id, "percentual": "100"}])],
        db=dados.db, usuario=dados.usuario)


def test_receita_direta_sem_cadastro_de_servico(dados):
    assert criar_receita(dados, dados.receita.id).status_code == 303
    item = dados.db.query(Atendimento).one()
    assert item.plano_conta_id == dados.receita.id
    assert item.descricao_servico == dados.receita.nome


@pytest.mark.parametrize("tipo,empresa,ativo", [("despesa", None, True), ("receita", 2, True), ("receita", None, False)])
def test_receita_rejeita_plano_incompativel(dados, tipo, empresa, ativo):
    plano = PlanoConta(tipo=tipo, cliente_id=empresa, ativo=ativo, nome="Invalido", grupo="Grupo", chave="invalido")
    dados.db.add(plano)
    dados.db.commit()
    criar_receita(dados, plano.id)
    assert dados.db.query(Atendimento).count() == 0


def test_filtros_receber_respeitam_empresa_status_forma_e_data(dados):
    for empresa, status, forma, dia in [(1, "pendente", "pix", 10), (1, "conciliado", "pix", 10),
                                      (2, "pendente", "pix", 10), (1, "pendente", "dinheiro", 10), (1, "pendente", "pix", 25)]:
        dados.db.add(Atendimento(condicao_pagamento="avista", cliente_id=empresa, status_conciliacao=status, forma_pagamento=forma,
            valor_servico=100, data_atendimento=date(2026, 1, dia)))
    dados.db.commit()
    response = chamar(lancamentos.listar_lancamentos, request=request(), data_inicio="2026-01-01", data_fim="2026-01-20",
        status_conciliacao="pendente", forma_pagamento="pix", db=dados.db, usuario=dados.usuario)
    assert len(response.context["atendimentos"]) == 1
    html = response.body.decode()
    assert 'href="/lancamentos/novo"' in html
    assert 'href="/contas-pagar/novo"' in html
    assert 'action="/lancamentos" method="post"' not in html


def importar(dados, monkeypatch, caminho):
    async def upload(*args):
        return str(caminho), caminho.name, ".xlsx"
    monkeypatch.setattr(contas_pagar, "salvar_upload_temporario", upload)
    return chamar(contas_pagar.importar_contas, request=request(), cliente_id=1,
        arquivo=SimpleNamespace(filename=caminho.name), db=dados.db, usuario=dados.usuario)


def test_modelo_download_upload_recorrencia_e_competencia(dados, monkeypatch, tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    gerar(caminho)
    wb = load_workbook(caminho)
    linha = {"descricao": "Aluguel", "valor": 100, "vencimento": date(2026, 1, 31),
             "data_competencia": date(2025, 12, 31), "forma_pagamento": "pix", "plano_de_contas": "D1",
             "recorrente": "sim", "recorrencia_intervalo": "mensal", "recorrencia_qtd": 3}
    for index, cell in enumerate(wb.active[1], 1):
        wb.active.cell(2, index, linha.get(cell.value))
    wb.save(caminho)
    wb.close()
    response = importar(dados, monkeypatch, caminho)
    assert "3 conta(s)" in unquote_plus(response.headers["location"])
    contas = dados.db.query(ContaPagar).order_by(ContaPagar.id).all()
    assert [c.vencimento for c in contas] == [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)]
    assert [c.data_competencia for c in contas] == [date(2025, 12, 31), date(2026, 1, 31), date(2026, 2, 28)]
    assert all(c.recorrencia_grupo_id == contas[0].id and c.plano_conta_id == dados.despesa.id and c.forma_pagamento.value == "pix" for c in contas)


@pytest.mark.parametrize("campo,valor", [("forma_pagamento", "invalido"), ("valor", -10),
    ("plano_de_contas", "inexistente"), ("recorrencia_qtd", "1.5"), ("recorrencia_qtd", 61),
    ("data_competencia", "invalida")])
def test_importacao_invalida_nao_salva_parcialmente(dados, monkeypatch, tmp_path, campo, valor):
    caminho = tmp_path / "entrada.xlsx"
    valida = dict(descricao="Valida", valor=100, vencimento=date(2026, 1, 31), plano_de_contas="D1")
    invalida = {**valida, campo: valor}
    pd.DataFrame([valida, invalida]).to_excel(caminho, index=False)
    response = importar(dados, monkeypatch, caminho)
    assert "Linha 3" in unquote_plus(response.headers["location"])
    assert dados.db.query(ContaPagar).count() == 0


@pytest.mark.parametrize("regime", ["caixa", "competencia"])
def test_telas_gestao_renderizam_e_preservam_regime(dados, regime):
    for funcao, path in [(gestao.gestao_receitas, "/gestao/receitas"), (gestao.gestao_despesas, "/gestao/despesas"),
                          (gestao.gestao_dre, "/gestao/dre"), (gestao.gestao_dre_apresentacao, "/gestao/dre/apresentacao")]:
        response = chamar(funcao, request=request(path), cliente_id=1, mes=1, ano=2026,
            regime=regime, db=dados.db, usuario=dados.usuario)
        assert response.status_code == 200
        assert response.context["regime"] == regime
        if "dre" in path:
            assert f"regime={regime}" in response.body.decode()
