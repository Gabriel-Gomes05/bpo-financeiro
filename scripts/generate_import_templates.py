"""Gera os modelos de importacao com as colunas aceitas pelos importadores."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


COLUNAS = {
    "descricao": "Obrigatoria. Descricao da despesa.",
    "fornecedor": "Nome do fornecedor (opcional).",
    "valor": "Obrigatorio. Valor positivo de cada ocorrencia, em reais.",
    "vencimento": "Obrigatorio. Data no formato DD/MM/AAAA ou data do Excel.",
    "data_competencia": "Data de competencia. Em branco, usa o vencimento.",
    "forma_pagamento": "pix, transferencia, cartao_credito, cartao_debito, dinheiro, boleto, cheque ou debito_automatico.",
    "plano_de_contas": "Codigo ou nome de uma conta de despesa ativa e disponivel para a empresa.",
    "recorrente": "sim ou nao. Em branco, considera o intervalo ou tipo informado.",
    "recorrencia_intervalo": "mensal, semanal, quinzenal ou personalizado. Avanca vencimento e competencia.",
    "recorrencia_dias": "Numero inteiro de dias, obrigatorio para intervalo personalizado.",
    "recorrencia_qtd": "Quantidade total de ocorrencias (incluindo a primeira), de 1 a 60. Padrao: 1.",
    "tipo": "pontual ou fixa. Coluna opcional para compatibilidade com modelos anteriores.",
    "observacao": "Observacoes (opcional).",
}


def gerar(destino: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contas a pagar"
    ws.append(list(COLUNAS))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:M1001"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A72")
        ws.column_dimensions[cell.column_letter].width = max(22, len(cell.value) + 3)
    for row in ws.iter_rows(min_row=2, max_row=1001, min_col=3, max_col=5):
        row[0].number_format = '#,##0.00'
        row[1].number_format = row[2].number_format = "dd/mm/yyyy"
    for coluna, opcoes in {
        "F": "pix,transferencia,cartao_credito,cartao_debito,dinheiro,boleto,cheque,debito_automatico",
        "H": "sim,nao",
        "I": "mensal,semanal,quinzenal,personalizado",
        "L": "pontual,fixa",
    }.items():
        validacao = DataValidation(type="list", formula1=f'"{opcoes}"', allow_blank=True)
        validacao.showErrorMessage = True
        validacao.errorTitle = "Valor invalido"
        validacao.error = "Selecione uma opcao da lista."
        ws.add_data_validation(validacao)
        validacao.add(f"{coluna}2:{coluna}1001")
    instrucoes = wb.create_sheet("Instrucoes")
    instrucoes.append(["Coluna", "Preenchimento"])
    for nome, descricao in COLUNAS.items():
        instrucoes.append([nome, descricao])
    instrucoes.column_dimensions["A"].width = 26
    instrucoes.column_dimensions["B"].width = 115
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


BPO_RECEITA = [
    "Data do atendimento", "Paciente", "CPF paciente/responsável", "Médico", "Especialidade",
    "Tipo de serviço", "Descrição do serviço", "Valor do serviço", "Condições do pagamento",
    "Parcela", "Data prevista de recebimento", "Forma de pagamento", "Taxa do cartão",
    "4 últimos dígitos do cartão", "Bandeira do cartão", "Valor líquido", "Data do crédito",
    "Destino do crédito", "Observação", "Plano de contas",
]

VISAO_CONTAS_RECEBER = [
    "Nome do cliente", "Data de competência", "Descrição", "Valor original da parcela (R$)",
    "Data prevista", "Forma de recebimento", "Conta bancária", "Data do último pagamento",
    "Categoria 1", "Centro de custo 1", "Observação",
]


def _cabecalho(ws, linha: int) -> None:
    for cell in ws[linha]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A72")
        ws.column_dimensions[cell.column_letter].width = max(18, len(str(cell.value)) + 3)


def gerar_bpo_receitas(destino: Path) -> None:
    """Aba RECEITA com cabecalho na linha 6; taxa do cartao em fracao (0,0185 = 1,85%)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "RECEITA"
    ws["A1"] = "Planilha BPO atualizado - receitas"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Mantenha o cabecalho na linha 6 e preencha a partir da linha 7."
    ws["A3"] = "Taxa do cartao em percentual do Excel (ex.: 1,85%). Datas em DD/MM/AAAA."
    ws.append([])
    ws.append([])
    ws.append(BPO_RECEITA)
    _cabecalho(ws, 6)
    ws.freeze_panes = "A7"
    for row in ws.iter_rows(min_row=7, max_row=1006):
        for idx in (0, 10, 16):
            row[idx].number_format = "dd/mm/yyyy"
        row[7].number_format = row[15].number_format = "#,##0.00"
        row[12].number_format = "0.00%"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def gerar_visao_contas_receber(destino: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Visão Contas a Receber"
    ws.append(VISAO_CONTAS_RECEBER)
    _cabecalho(ws, 1)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2, max_row=1001):
        for idx in (1, 4, 7):
            row[idx].number_format = "dd/mm/yyyy"
        row[3].number_format = "#,##0.00"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


if __name__ == "__main__":
    modelos = Path(__file__).resolve().parents[1] / "static/modelos"
    gerar(modelos / "modelo_contas_pagar.xlsx")
    gerar_bpo_receitas(modelos / "modelo_bpo_receitas.xlsx")
    gerar_visao_contas_receber(modelos / "modelo_visao_contas_receber.xlsx")
